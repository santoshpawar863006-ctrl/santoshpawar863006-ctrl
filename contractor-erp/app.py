import sys
from pathlib import Path

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QStackedWidget, QFrame, QTableWidget, QTableWidgetItem, QHeaderView, QDialog,
    QFormLayout, QLineEdit, QDoubleSpinBox, QComboBox, QTextEdit, QDialogButtonBox,
    QMessageBox, QFileDialog, QGridLayout
)

from db import init_db, query, execute, scalar, now, backup_database, DB_PATH

APP_STYLE = """
QMainWindow { background: #f4f6f8; }
QWidget { font-family: Segoe UI, Arial; font-size: 13px; color: #202124; }
#Sidebar { background: #121821; border: 0; }
#Brand { color: white; font-size: 20px; font-weight: 700; padding: 18px 12px; }
QPushButton#NavButton { color: #d8dee9; background: transparent; text-align: left; padding: 11px 16px; border: none; border-radius: 6px; }
QPushButton#NavButton:hover { background: #202936; color: white; }
QPushButton#NavButton:checked { background: #2f6fed; color: white; }
QPushButton#Primary { background: #2f6fed; color: white; border: none; border-radius: 7px; padding: 9px 14px; font-weight: 600; }
QPushButton#Danger { background: #b3261e; color: white; border: none; border-radius: 7px; padding: 8px 12px; }
QPushButton#Secondary { background: white; border: 1px solid #d0d7de; border-radius: 7px; padding: 8px 12px; }
QFrame#Card { background: white; border: 1px solid #e1e5ea; border-radius: 10px; }
QLabel#CardTitle { color: #687078; font-size: 12px; }
QLabel#CardValue { font-size: 24px; font-weight: 700; color: #111827; }
QLabel#PageTitle { font-size: 23px; font-weight: 700; }
QTableWidget { background: white; border: 1px solid #e1e5ea; border-radius: 8px; gridline-color: #eef1f4; selection-background-color: #dce7ff; }
QHeaderView::section { background: #f8fafc; padding: 8px; border: 0; border-bottom: 1px solid #e5e7eb; font-weight: 600; }
QLineEdit, QComboBox, QDoubleSpinBox, QTextEdit { background: white; border: 1px solid #cfd6dd; border-radius: 6px; padding: 7px; }
"""

def money(v):
    try:
        return f"₹{float(v):,.2f}"
    except Exception:
        return "₹0.00"

class Card(QFrame):
    def __init__(self, title, value="0"):
        super().__init__()
        self.setObjectName("Card")
        layout = QVBoxLayout(self)
        layout.setContentsMargins(18, 16, 18, 16)
        t = QLabel(title); t.setObjectName("CardTitle")
        self.value = QLabel(value); self.value.setObjectName("CardValue")
        layout.addWidget(t); layout.addWidget(self.value)

class Dashboard(QWidget):
    def __init__(self):
        super().__init__()
        root = QVBoxLayout(self); root.setContentsMargins(24, 22, 24, 24); root.setSpacing(16)
        title = QLabel("Dashboard"); title.setObjectName("PageTitle"); root.addWidget(title)
        grid = QGridLayout(); grid.setSpacing(14)
        self.cards = {
            "projects": Card("Active Projects"), "contract": Card("Total Contract Value"),
            "received": Card("Client Amount Received"), "outstanding": Card("Client Outstanding"),
            "expenses": Card("Recorded Expenses"), "materials": Card("Material Purchases"),
            "labour": Card("Labour Cost"), "profit": Card("Estimated Current Profit"),
        }
        for i, card in enumerate(self.cards.values()): grid.addWidget(card, i // 4, i % 4)
        root.addLayout(grid)
        section = QLabel("Recent Projects"); section.setStyleSheet("font-size:16px;font-weight:700;margin-top:8px;"); root.addWidget(section)
        self.table = QTableWidget(0, 6)
        self.table.setHorizontalHeaderLabels(["Project", "Client", "Location", "Contract Value", "Progress", "Status"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch); self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        root.addWidget(self.table, 1)
    def refresh(self):
        active = scalar("SELECT COUNT(*) FROM projects WHERE status='Active'")
        contract = scalar("SELECT COALESCE(SUM(contract_value),0) FROM projects")
        billed = scalar("SELECT COALESCE(SUM(amount),0) FROM bills")
        received = scalar("SELECT COALESCE(SUM(received),0) FROM bills")
        exp = scalar("SELECT COALESCE(SUM(amount),0) FROM expenses")
        mat = scalar("SELECT COALESCE(SUM(amount),0) FROM materials")
        labour = scalar("SELECT COALESCE(SUM(attendance*daily_wage),0) FROM labour")
        cost = exp + mat + labour; profit = received - cost
        self.cards["projects"].value.setText(str(active)); self.cards["contract"].value.setText(money(contract))
        self.cards["received"].value.setText(money(received)); self.cards["outstanding"].value.setText(money(max(billed-received,0)))
        self.cards["expenses"].value.setText(money(exp)); self.cards["materials"].value.setText(money(mat)); self.cards["labour"].value.setText(money(labour)); self.cards["profit"].value.setText(money(profit))
        rows = query("SELECT name, client_name, location, contract_value, progress, status FROM projects ORDER BY id DESC LIMIT 8")
        self.table.setRowCount(len(rows))
        for r,row in enumerate(rows):
            vals=[row['name'],row['client_name'] or '',row['location'] or '',money(row['contract_value']),f"{row['progress']:.0f}%",row['status']]
            for c,val in enumerate(vals): self.table.setItem(r,c,QTableWidgetItem(str(val)))

class GenericDialog(QDialog):
    def __init__(self, title, fields, parent=None):
        super().__init__(parent); self.setWindowTitle(title); self.resize(460, 540); self.widgets={}; form=QFormLayout(self)
        for name,label,kind,options in fields:
            if kind=='text': w=QLineEdit()
            elif kind=='number':
                w=QDoubleSpinBox(); w.setRange(0,1_000_000_000); w.setDecimals(2)
            elif kind=='combo': w=QComboBox(); w.addItems(options or [])
            elif kind=='project':
                w=QComboBox(); w.addItem('Select project...',None)
                for p in query('SELECT id,name FROM projects ORDER BY name'): w.addItem(p['name'],p['id'])
                w.setProperty('returnData',True)
            elif kind=='multiline': w=QTextEdit(); w.setFixedHeight(90)
            else: w=QLineEdit()
            self.widgets[name]=w; form.addRow(label,w)
        buttons=QDialogButtonBox(QDialogButtonBox.Save|QDialogButtonBox.Cancel); buttons.accepted.connect(self.accept); buttons.rejected.connect(self.reject); form.addRow(buttons)
    def data(self):
        out={}
        for name,w in self.widgets.items():
            if isinstance(w,QDoubleSpinBox): out[name]=w.value()
            elif isinstance(w,QComboBox): out[name]=w.currentData() if w.property('returnData') else w.currentText()
            elif isinstance(w,QTextEdit): out[name]=w.toPlainText().strip()
            else: out[name]=w.text().strip()
        return out

class CrudPage(QWidget):
    def __init__(self,title,table,columns,fields,insert_sql,select_sql,delete_sql=None):
        super().__init__(); self.title=title; self.table_name=table; self.columns=columns; self.fields=fields; self.insert_sql=insert_sql; self.select_sql=select_sql; self.delete_sql=delete_sql or f'DELETE FROM {table} WHERE id=?'
        root=QVBoxLayout(self); root.setContentsMargins(24,22,24,24); root.setSpacing(14)
        header=QHBoxLayout(); t=QLabel(title); t.setObjectName('PageTitle'); header.addWidget(t); header.addStretch()
        add=QPushButton('+ Add'); add.setObjectName('Primary'); add.clicked.connect(self.add_record); header.addWidget(add)
        refresh=QPushButton('Refresh'); refresh.setObjectName('Secondary'); refresh.clicked.connect(self.refresh); header.addWidget(refresh)
        delete=QPushButton('Delete Selected'); delete.setObjectName('Danger'); delete.clicked.connect(self.delete_selected); header.addWidget(delete); root.addLayout(header)
        self.table=QTableWidget(0,len(columns)); self.table.setHorizontalHeaderLabels([c[0] for c in columns]); self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch); self.table.setSelectionBehavior(QTableWidget.SelectRows); self.table.setEditTriggers(QTableWidget.NoEditTriggers); root.addWidget(self.table,1)
    def add_record(self):
        dlg=GenericDialog(f"Add {self.title.rstrip('s')}",self.fields,self)
        if dlg.exec()!=QDialog.Accepted: return
        data=dlg.data()
        try:
            execute(self.insert_sql,tuple(data[k] for k,*_ in self.fields)+(now(),)); self.refresh()
        except Exception as e: QMessageBox.critical(self,'Could not save',str(e))
    def delete_selected(self):
        row=self.table.currentRow()
        if row<0: QMessageBox.information(self,'Select a row','Please select a record to delete.'); return
        record_id=self.table.item(row,0).data(Qt.UserRole)
        if QMessageBox.question(self,'Delete','Delete the selected record?')!=QMessageBox.Yes: return
        execute(self.delete_sql,(record_id,)); self.refresh()
    def refresh(self):
        rows=query(self.select_sql); self.table.setRowCount(len(rows))
        for r,row in enumerate(rows):
            for c,(label,key,fmt) in enumerate(self.columns):
                val=row[key]
                if fmt=='money': val=money(val)
                elif fmt=='percent': val=f"{float(val or 0):.0f}%"
                item=QTableWidgetItem('' if val is None else str(val))
                if c==0: item.setData(Qt.UserRole,row['id'])
                self.table.setItem(r,c,item)

class ReportsPage(QWidget):
    def __init__(self):
        super().__init__(); root=QVBoxLayout(self); root.setContentsMargins(24,22,24,24)
        title=QLabel('Reports'); title.setObjectName('PageTitle'); root.addWidget(title)
        self.summary=QLabel(); self.summary.setWordWrap(True); self.summary.setStyleSheet('background:white;border:1px solid #e1e5ea;border-radius:10px;padding:18px;font-size:14px;'); root.addWidget(self.summary)
        actions=QHBoxLayout();
        excel=QPushButton('Export Summary to Excel'); excel.setObjectName('Primary'); excel.clicked.connect(self.export_excel); actions.addWidget(excel)
        pdf=QPushButton('Export Summary to PDF'); pdf.setObjectName('Primary'); pdf.clicked.connect(self.export_pdf); actions.addWidget(pdf)
        backup=QPushButton('Create Database Backup'); backup.setObjectName('Secondary'); backup.clicked.connect(self.do_backup); actions.addWidget(backup); actions.addStretch(); root.addLayout(actions); root.addStretch()
    def get_summary_data(self):
        projects=scalar('SELECT COUNT(*) FROM projects'); contract=scalar('SELECT COALESCE(SUM(contract_value),0) FROM projects'); received=scalar('SELECT COALESCE(SUM(received),0) FROM bills'); billed=scalar('SELECT COALESCE(SUM(amount),0) FROM bills'); materials=scalar('SELECT COALESCE(SUM(amount),0) FROM materials'); labour=scalar('SELECT COALESCE(SUM(attendance*daily_wage),0) FROM labour'); expenses=scalar('SELECT COALESCE(SUM(amount),0) FROM expenses'); cost=materials+labour+expenses
        return [('Projects',projects),('Contract Value',contract),('Total Billed',billed),('Client Receipts',received),('Client Outstanding',max(billed-received,0)),('Material Purchases',materials),('Labour Cost',labour),('Other Expenses',expenses),('Recorded Cost',cost),('Current Cash Profit',received-cost)]
    def refresh(self):
        d=dict(self.get_summary_data()); self.summary.setText(f"<b>Company Summary</b><br><br>Projects: <b>{d['Projects']}</b><br>Contract Value: <b>{money(d['Contract Value'])}</b><br>Client Receipts: <b>{money(d['Client Receipts'])}</b><br>Material Purchases: <b>{money(d['Material Purchases'])}</b><br>Labour Cost: <b>{money(d['Labour Cost'])}</b><br>Other Expenses: <b>{money(d['Other Expenses'])}</b><br>Recorded Cost: <b>{money(d['Recorded Cost'])}</b><br>Current Cash Profit: <b>{money(d['Current Cash Profit'])}</b><br><br>Database location:<br>{DB_PATH}")
    def export_excel(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font
        path,_=QFileDialog.getSaveFileName(self,'Export Excel','Contractor_ERP_Summary.xlsx','Excel (*.xlsx)')
        if not path:return
        if not path.lower().endswith('.xlsx'): path+='.xlsx'
        wb=Workbook(); ws=wb.active; ws.title='Company Summary'; ws['A1']='CONTRACTOR ERP - COMPANY SUMMARY'; ws['A1'].font=Font(size=16,bold=True); ws.append(['Metric','Value'])
        for metric,value in self.get_summary_data(): ws.append([metric,value])
        ws.column_dimensions['A'].width=28; ws.column_dimensions['B'].width=20; wb.save(path); QMessageBox.information(self,'Excel exported',f'Saved to:\n{path}')
    def export_pdf(self):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        path,_=QFileDialog.getSaveFileName(self,'Export PDF','Contractor_ERP_Summary.pdf','PDF (*.pdf)')
        if not path:return
        if not path.lower().endswith('.pdf'): path+='.pdf'
        c=canvas.Canvas(path,pagesize=A4); width,height=A4; y=height-55; c.setFont('Helvetica-Bold',16); c.drawString(45,y,'CONTRACTOR ERP - COMPANY SUMMARY'); y-=35; c.setFont('Helvetica',11)
        for metric,value in self.get_summary_data():
            display=str(value) if metric=='Projects' else f'INR {float(value):,.2f}'; c.drawString(50,y,metric); c.drawRightString(width-50,y,display); y-=24
        c.save(); QMessageBox.information(self,'PDF exported',f'Saved to:\n{path}')
    def do_backup(self):
        path,_=QFileDialog.getSaveFileName(self,'Save Backup','contractor_erp_backup.db','Database (*.db)')
        if not path:return
        backup_database(path); QMessageBox.information(self,'Backup complete',f'Backup created:\n{path}')

def project_fields():
    return [('name','Project Name','text',None),('client_name','Client Name','text',None),('location','Location','text',None),('start_date','Start Date (YYYY-MM-DD)','text',None),('end_date','End Date (YYYY-MM-DD)','text',None),('contract_value','Contract Value','number',None),('status','Status','combo',['Active','On Hold','Completed','Cancelled']),('progress','Progress %','number',None),('notes','Notes','multiline',None)]

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__(); self.setWindowTitle('Contractor ERP'); self.resize(1440,860)
        root=QWidget(); self.setCentralWidget(root); layout=QHBoxLayout(root); layout.setContentsMargins(0,0,0,0); layout.setSpacing(0)
        sidebar=QFrame(); sidebar.setObjectName('Sidebar'); sidebar.setFixedWidth(220); s=QVBoxLayout(sidebar); s.setContentsMargins(10,8,10,12)
        brand=QLabel('CONTRACTOR ERP'); brand.setObjectName('Brand'); s.addWidget(brand)
        self.stack=QStackedWidget(); self.pages=[]; self.nav_buttons=[]
        self.dashboard=Dashboard(); self.add_page(s,'Dashboard',self.dashboard)
        self.add_page(s,'Projects',CrudPage('Projects','projects',[('Project','name',None),('Client','client_name',None),('Location','location',None),('Contract Value','contract_value','money'),('Progress','progress','percent'),('Status','status',None)],project_fields(),'INSERT INTO projects (name,client_name,location,start_date,end_date,contract_value,status,progress,notes,created_at) VALUES (?,?,?,?,?,?,?,?,?,?)','SELECT * FROM projects ORDER BY id DESC'))
        self.add_page(s,'Clients',CrudPage('Clients','clients',[('Name','name',None),('Phone','phone',None),('Email','email',None),('GSTIN','gstin',None),('Address','address',None)],[('name','Name','text',None),('phone','Phone','text',None),('email','Email','text',None),('address','Address','multiline',None),('gstin','GSTIN','text',None),('notes','Notes','multiline',None)],'INSERT INTO clients (name,phone,email,address,gstin,notes,created_at) VALUES (?,?,?,?,?,?,?)','SELECT * FROM clients ORDER BY id DESC'))
        self.add_page(s,'BOQ / Estimates',CrudPage('BOQ Items','boq',[('Project ID','project_id',None),('Item','item_code',None),('Description','description',None),('Unit','unit',None),('Qty','quantity',None),('Rate','rate','money')],[('project_id','Project','project',None),('item_code','Item Code','text',None),('description','Description','text',None),('unit','Unit','text',None),('quantity','Quantity','number',None),('rate','Rate','number',None),('material_rate','Material Rate','number',None),('labour_rate','Labour Rate','number',None)],'INSERT INTO boq (project_id,item_code,description,unit,quantity,rate,material_rate,labour_rate,created_at) VALUES (?,?,?,?,?,?,?,?,?)','SELECT * FROM boq ORDER BY id DESC'))
        self.add_page(s,'Materials',CrudPage('Materials','materials',[('Date','date',None),('Project ID','project_id',None),('Material','material',None),('Supplier','supplier',None),('Qty','quantity',None),('Rate','rate','money'),('Amount','amount','money')],[('project_id','Project','project',None),('date','Date','text',None),('supplier','Supplier','text',None),('material','Material','text',None),('unit','Unit','text',None),('quantity','Quantity','number',None),('rate','Rate','number',None),('amount','Amount','number',None),('invoice_no','Invoice No.','text',None),('notes','Notes','multiline',None)],'INSERT INTO materials (project_id,date,supplier,material,unit,quantity,rate,amount,invoice_no,notes,created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?)','SELECT * FROM materials ORDER BY id DESC'))
        self.add_page(s,'Labour',CrudPage('Labour','labour',[('Date','date',None),('Project ID','project_id',None),('Worker','worker_name',None),('Trade','trade',None),('Attendance','attendance',None),('Daily Wage','daily_wage','money'),('Paid','paid','money')],[('project_id','Project','project',None),('date','Date','text',None),('worker_name','Worker Name','text',None),('trade','Trade','text',None),('attendance','Attendance Days','number',None),('daily_wage','Daily Wage','number',None),('advance','Advance','number',None),('paid','Paid','number',None),('notes','Notes','multiline',None)],'INSERT INTO labour (project_id,date,worker_name,trade,attendance,daily_wage,advance,paid,notes,created_at) VALUES (?,?,?,?,?,?,?,?,?,?)','SELECT * FROM labour ORDER BY id DESC'))
        self.add_page(s,'Expenses',CrudPage('Expenses','expenses',[('Date','date',None),('Project ID','project_id',None),('Category','category',None),('Description','description',None),('Amount','amount','money'),('Mode','payment_mode',None)],[('project_id','Project','project',None),('date','Date','text',None),('category','Category','combo',['Transport','Fuel','Machinery','Office','Site','Food','Other']),('description','Description','text',None),('amount','Amount','number',None),('payment_mode','Payment Mode','combo',['Cash','UPI','Bank','Cheque','Other']),('reference','Reference','text',None)],'INSERT INTO expenses (project_id,date,category,description,amount,payment_mode,reference,created_at) VALUES (?,?,?,?,?,?,?,?)','SELECT * FROM expenses ORDER BY id DESC'))
        self.add_page(s,'Suppliers',CrudPage('Suppliers','suppliers',[('Name','name',None),('Phone','phone',None),('Email','email',None),('GSTIN','gstin',None),('Address','address',None)],[('name','Name','text',None),('phone','Phone','text',None),('email','Email','text',None),('address','Address','multiline',None),('gstin','GSTIN','text',None),('notes','Notes','multiline',None)],'INSERT INTO suppliers (name,phone,email,address,gstin,notes,created_at) VALUES (?,?,?,?,?,?,?)','SELECT * FROM suppliers ORDER BY id DESC'))
        self.add_page(s,'Billing / Payments',CrudPage('Bills','bills',[('Bill No.','bill_no',None),('Project ID','project_id',None),('Date','bill_date',None),('Amount','amount','money'),('Received','received','money'),('Due Date','due_date',None),('Status','status',None)],[('project_id','Project','project',None),('bill_no','Bill No.','text',None),('bill_date','Bill Date','text',None),('description','Description','text',None),('amount','Bill Amount','number',None),('received','Received','number',None),('due_date','Due Date','text',None),('status','Status','combo',['Pending','Part Paid','Paid','Overdue'])],'INSERT INTO bills (project_id,bill_no,bill_date,description,amount,received,due_date,status,created_at) VALUES (?,?,?,?,?,?,?,?,?)','SELECT * FROM bills ORDER BY id DESC'))
        self.reports=ReportsPage(); self.add_page(s,'Reports / Backup',self.reports); s.addStretch(); footer=QLabel('Local database\nStored on this laptop'); footer.setStyleSheet('color:#98a2b3;padding:12px;'); s.addWidget(footer)
        layout.addWidget(sidebar); layout.addWidget(self.stack,1)
        if self.nav_buttons:self.nav_buttons[0].setChecked(True)
        self.switch_page(0)
    def add_page(self,sidebar_layout,name,page):
        idx=self.stack.addWidget(page); self.pages.append(page); btn=QPushButton(name); btn.setObjectName('NavButton'); btn.setCheckable(True); btn.clicked.connect(lambda checked=False,i=idx:self.switch_page(i)); sidebar_layout.addWidget(btn); self.nav_buttons.append(btn)
    def switch_page(self,idx):
        self.stack.setCurrentIndex(idx)
        for i,b in enumerate(self.nav_buttons): b.setChecked(i==idx)
        page=self.pages[idx]
        if hasattr(page,'refresh'): page.refresh()

def main():
    init_db(); app=QApplication(sys.argv); app.setApplicationName('Contractor ERP'); app.setStyleSheet(APP_STYLE); window=MainWindow(); window.show(); sys.exit(app.exec())

if __name__=='__main__': main()
