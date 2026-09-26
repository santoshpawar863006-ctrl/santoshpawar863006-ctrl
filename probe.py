import io, requests
from openpyxl import load_workbook
from collect_results import parse_statement
API = "https://kppp.karnataka.gov.in/supplier-registration-service/v1/api/portal-service"
H = {"Accept": "*/*", "Origin": "https://kppp.karnataka.gov.in", "Referer": "https://kppp.karnataka.gov.in/",
     "Post": "CONTRACTOR-EPROC-CONTRACTOR", "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/124.0 Safari/537.36"}
for nit in ("54429", "57506"):
    r = requests.get(f"{API}/tender-eval/{nit}/commercial-evaluation/tender-category/GOODS/commercial-comparison/download-detailed", headers=H, timeout=60)
    print("==", nit, r.status_code, len(r.content))
    if r.content[:2] != b"PK": print(r.text[:300]); continue
    for ws in load_workbook(io.BytesIO(r.content)).worksheets:
        print("sheet", ws.title)
        for row in ws.iter_rows(values_only=True):
            if any(v not in (None, "") for v in row): print("  ", [v for v in row][:24])
    try: print("parsed", parse_statement(r.content))
    except Exception as e: print("parse failed", repr(e))
